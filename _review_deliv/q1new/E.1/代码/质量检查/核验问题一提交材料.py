"""Read-only audit of Problem 1 inputs, features and paper; write delivery evidence.

This script never rewrites the official videos, labels or feature arrays.
Run from any working directory with Python, NumPy, openpyxl and python-docx.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import wave
import zipfile
from collections import defaultdict
from pathlib import Path

import numpy as np
import openpyxl
from docx import Document


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "交付候选"
MANIFEST = ROOT / "输入/样本清单/样本清单.csv"
LABEL = ROOT / "输入/原始标签/label-100.xlsx"
WORD_INDEX = ROOT / "结果/08_原文词级跨模态对齐/逐词对齐索引.csv"
TEXT_MATRIX = ROOT / "结果/05_BERT_base_uncased/原文逐词BERT特征.npy"
PAPER = ROOT / "论文/问题一_五节结构与图表精修稿_图2优化版.docx"
FIGURE = ROOT / "论文/图表精修/论文插图/fig01_flow.png"
STATS = ROOT / "结果/10_全量质量检查与论文图表/全量样本统计.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def check(condition: bool, message: str, failures: list[str]) -> None:
    if not condition:
        failures.append(message)


def rel(path: Path) -> str:
    return "E.1/" + path.relative_to(ROOT).as_posix()


def main() -> None:
    OUT.mkdir(exist_ok=True)
    failures: list[str] = []
    manifest = read_csv(MANIFEST)
    index = read_csv(WORD_INDEX)
    stats = json.loads(STATS.read_text(encoding="utf-8"))
    ids = [row["sample_id"] for row in manifest]
    id_set = set(ids)
    check(len(manifest) == len(id_set) == 100, "样本清单不是100个唯一编号", failures)
    check(len(stats) == 100 and {r["sample_id"] for r in stats} == id_set,
          "全量统计与样本清单编号不一致", failures)
    stats_by_id = {r["sample_id"]: r for r in stats}

    official_label = Path(manifest[0]["source_video_path"]).parent.parent / "label-100.xlsx"
    check(official_label.is_file() and sha256(official_label) == sha256(LABEL),
          "原始标签与E.1副本SHA-256不一致", failures)
    workbook = openpyxl.load_workbook(LABEL, read_only=True, data_only=True)
    sheet_rows = list(workbook["label"].values)
    workbook.close()
    check(tuple(sheet_rows[0][:5]) == ("video_id", "clip_id", "text", "label", "annotation")
          and len(sheet_rows) == 101, "原始标签不是预期的100行和5列", failures)
    for row in manifest:
        excel_row = sheet_rows[int(row["source_excel_row"]) - 1]
        for field, value in zip(("video_id", "clip_id", "text", "label_raw", "annotation"), excel_row[:5]):
            check(row[field] == str(value), f'{row["sample_id"]}: {field}与原始Excel不一致', failures)
        check(row["sample_id"] == f'{row["video_id"]}$_${row["clip_id"]}',
              f'{row["sample_id"]}: 编号构造错误', failures)

    by_id: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in index:
        by_id[row["sample_id"]].append(row)
    check(len(index) == 1934 and set(by_id) == id_set,
          "逐词索引不是100样本、1934词", failures)
    global_text = np.load(TEXT_MATRIX, mmap_mode="r", allow_pickle=False)
    check(global_text.shape == (1934, 768) and global_text.dtype == np.float32,
          "BERT全量原文词矩阵维度或类型错误", failures)

    expected_dirs = {
        "audio": ROOT / "结果/06_Librosa74/逐样本",
        "visual": ROOT / "结果/07_FaceLandmarker52/逐样本",
        "aligned": ROOT / "结果/08_原文词级跨模态对齐/逐样本",
        "asr": ROOT / "结果/03_Whisper_small_en/逐样本",
    }
    for kind, folder in expected_dirs.items():
        ext = ".json" if kind == "asr" else ".npz"
        found = list(folder.glob(f"*{ext}"))
        check(len(found) == 100, f"{kind}逐样本文件不是100份", failures)
        if kind != "asr":
            check({p.stem for p in found} == id_set, f"{kind}文件编号不一致", failures)

    totals = defaultdict(int)
    delivery_rows: list[dict[str, str | int | float]] = []
    no_time_ids: list[str] = []
    no_visual_ids: list[str] = []
    for order, item in enumerate(manifest, 1):
        sid = item["sample_id"]
        wav = ROOT / "输入/音频WAV" / f"{sid}.wav"
        video = ROOT / "输入/原始视频" / item["video_relative_path"]
        source_video = Path(item["source_video_path"])
        a_path = expected_dirs["audio"] / f"{sid}.npz"
        v_path = expected_dirs["visual"] / f"{sid}.npz"
        main_path = expected_dirs["aligned"] / f"{sid}.npz"
        needed = [wav, video, source_video, a_path, v_path, main_path]
        if not all(p.is_file() for p in needed):
            failures.append(f"{sid}: 原始媒体或输出文件缺失")
            continue
        video_hash = sha256(video)
        check(video_hash == item["sha256"] == sha256(source_video),
              f"{sid}: 原始MP4、副本与清单哈希不一致", failures)
        with wave.open(str(wav), "rb") as wave_file:
            check(wave_file.getframerate() == 16000 and wave_file.getnchannels() == 1,
                  f"{sid}: WAV采样率或声道不符", failures)
        word_rows = by_id[sid]
        n = len(word_rows)
        check([int(r["official_word_index"]) for r in word_rows] == list(range(1, n + 1)),
              f"{sid}: 原文词序不连续", failures)
        text_rows = [int(r["text_feature_row_0based"]) for r in word_rows]
        check(text_rows == list(range(text_rows[0], text_rows[0] + n)),
              f"{sid}: BERT全局行号不连续", failures)

        with (np.load(a_path, allow_pickle=False) as audio,
              np.load(v_path, allow_pickle=False) as visual,
              np.load(main_path, allow_pickle=False) as main):
            a_raw = audio["features"]
            v_raw = visual["features"]
            t = main["text"]
            a = main["audio"]
            v = main["visual"]
            check(a_raw.ndim == 2 and a_raw.shape[1] == 74 and a_raw.dtype == np.float32,
                  f"{sid}: 原始声学特征维度错误", failures)
            check(v_raw.ndim == 2 and v_raw.shape[1] == 52 and v_raw.dtype == np.float32,
                  f"{sid}: 原始视觉特征维度错误", failures)
            check((t.shape, a.shape, v.shape) == ((n, 768), (n, 74), (n, 52)),
                  f"{sid}: 主矩阵行数或维度错误", failures)
            check(all(x.dtype == np.float32 for x in (t, a, v)),
                  f"{sid}: 主矩阵类型错误", failures)
            check(int(main["valid_length"]) == n and np.all(main["text_valid_mask"] == 1),
                  f"{sid}: 有效长度或文本掩码错误", failures)
            check(np.array_equal(global_text[text_rows], t),
                  f"{sid}: 主文本向量与原始BERT矩阵不一致", failures)
            tm = main["word_time_valid_mask"].astype(bool)
            am = main["audio_valid_mask"].astype(bool)
            vm = main["visual_valid_mask"].astype(bool)
            fm = main["audio_f0_valid_mask"].astype(bool)
            starts = main["word_start_wav_s"]
            ends = main["word_end_wav_s"]
            check(all(len(main[k]) == n for k in (
                "word_start_wav_s", "word_end_wav_s", "word_start_source_s",
                "word_end_source_s", "text_valid_mask", "word_time_valid_mask",
                "audio_valid_mask", "audio_f0_valid_mask", "visual_valid_mask",
                "audio_frame_count", "visual_frame_count")),
                f"{sid}: 主时间/掩码/帧数字段长度错误", failures)
            check(np.all(np.isfinite(starts[tm])) and np.all(np.isfinite(ends[tm]))
                  and np.all(starts[tm] < ends[tm]) and np.all(starts[tm] >= 0),
                  f"{sid}: 有效词时间范围错误", failures)
            check(np.all(np.isnan(starts[~tm])) and np.all(np.isnan(ends[~tm])),
                  f"{sid}: 无时间词未使用NaN", failures)
            check(np.all(am <= tm) and np.all(vm <= tm) and np.all(fm <= am),
                  f"{sid}: 模态掩码与词时间不相容", failures)
            check(np.count_nonzero(a[~am]) == 0 and np.count_nonzero(v[~vm]) == 0,
                  f"{sid}: 缺失模态未使用零占位", failures)
            check(np.all(np.diff(audio["frame_time_wav_s"]) >= 0)
                  and np.all(np.diff(visual["frame_time_source_s"]) >= 0),
                  f"{sid}: 原始帧时间顺序错误", failures)
            check(len(audio["frame_time_wav_s"]) == len(a_raw)
                  and len(visual["frame_time_source_s"]) == len(v_raw),
                  f"{sid}: 原始帧时间与特征行不匹配", failures)
            for key, mask in (("word_time_valid_mask", tm), ("audio_valid_mask", am),
                              ("visual_valid_mask", vm)):
                check(np.array_equal(mask.astype(np.uint8),
                                     np.array([int(r[key]) for r in word_rows], dtype=np.uint8)),
                      f"{sid}: 主NPZ与逐词CSV的{key}不一致", failures)
            check(int(tm.sum()) == stats_by_id[sid]["timed_words"]
                  and int(am.sum()) == stats_by_id[sid]["audio_words"]
                  and int(vm.sum()) == stats_by_id[sid]["visual_words"],
                  f"{sid}: 主NPZ与全量统计的词数不一致", failures)
            totals["words"] += n
            totals["timed_words"] += int(tm.sum())
            totals["audio_words"] += int(am.sum())
            totals["visual_words"] += int(vm.sum())
            totals["audio_frames"] += len(a_raw)
            totals["sampled_video_frames"] += len(v_raw)
            totals["valid_face_frames"] += int(visual["face_valid_mask"].sum())
            if not tm.any():
                no_time_ids.append(sid)
            if not vm.any():
                no_visual_ids.append(sid)

        delivery_rows.append({
            "序号": order,
            "sample_id": sid,
            "原始视频相对路径": item["video_relative_path"],
            "原始Excel行号": item["source_excel_row"],
            "原始视频SHA256": video_hash,
            "视频有效时长_s": stats_by_id[sid]["video_duration_s"],
            "WAV有效时长_s": stats_by_id[sid]["wav_duration_s"],
            "原文词数_有效长度": n,
            "主时间有效词数": stats_by_id[sid]["timed_words"],
            "声学有效词数": stats_by_id[sid]["audio_words"],
            "视觉有效词数": stats_by_id[sid]["visual_words"],
            "声学原始帧数": stats_by_id[sid]["audio_frames"],
            "视觉采样帧数": stats_by_id[sid]["visual_sampled_frames"],
            "视觉有效帧数": stats_by_id[sid]["visual_valid_frames"],
            "文本维度": 768, "声学维度": 74, "视觉维度": 52,
            "对齐粒度": "原文词",
            "文本全量矩阵": rel(TEXT_MATRIX),
            "文本全量矩阵起始行_0based": text_rows[0],
            "文本全量矩阵结束行_exclusive": text_rows[-1] + 1,
            "声学逐帧文件": rel(a_path),
            "声学逐帧SHA256": sha256(a_path),
            "视觉逐帧文件": rel(v_path),
            "视觉逐帧SHA256": sha256(v_path),
            "词级三模态文件": rel(main_path),
            "词级三模态SHA256": sha256(main_path),
        })

    for key, expected in {
        "words": 1934, "timed_words": 1421, "audio_words": 1421,
        "visual_words": 1228, "audio_frames": 77755,
        "sampled_video_frames": 7889, "valid_face_frames": 6471,
    }.items():
        check(totals[key] == expected, f"全量{key}={totals[key]}，预期{expected}", failures)
    check(len(no_time_ids) == 16 and len(no_visual_ids) == 21,
          "无主时间或无主视觉样本数量与论文不一致", failures)

    paper = Document(PAPER)
    appendix_rows = [row for table in paper.tables for row in table.rows[1:]
                     if len(table.columns) == 9]
    paper_ids = [row.cells[1].text.strip() for row in appendix_rows]
    check(len(paper_ids) == 100 and set(paper_ids) == id_set,
          "论文附表A1未覆盖完整100条样本", failures)
    captions = [p.text for p in paper.paragraphs if p.text.startswith("图 ")]
    check(all(any(c.startswith(f"图 {i} ") for c in captions) for i in range(2, 10)),
          "论文图2—图9缺失", failures)
    with zipfile.ZipFile(PAPER) as archive:
        check(archive.read("word/media/image1.png") == FIGURE.read_bytes(),
              "论文图2与定稿流程图不一致", failures)

    index_file = OUT / "100条交付索引.csv"
    if len(delivery_rows) == 100:
        with index_file.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(delivery_rows[0]), lineterminator="\n")
            writer.writeheader()
            writer.writerows(delivery_rows)
    report = {
        "status": "pass" if not failures else "fail",
        "source_label_sha256": sha256(LABEL),
        "paper_file": rel(PAPER),
        "paper_appendix_sample_rows": len(paper_ids),
        "paper_figures_2_to_9_present": all(any(c.startswith(f"图 {i} ") for c in captions)
                                             for i in range(2, 10)),
        "sample_count": len(manifest),
        "raw_mp4_count": len(list((ROOT / "输入/原始视频").rglob("*.mp4"))),
        "raw_wav_count": len(list((ROOT / "输入/音频WAV").rglob("*.wav"))),
        "raw_audio_npz_count": len(list(expected_dirs["audio"].glob("*.npz"))),
        "raw_visual_npz_count": len(list(expected_dirs["visual"].glob("*.npz"))),
        "aligned_npz_count": len(list(expected_dirs["aligned"].glob("*.npz"))),
        "word_index_rows": len(index),
        "totals": dict(totals),
        "samples_without_primary_time": no_time_ids,
        "samples_without_primary_visual": no_visual_ids,
        "delivery_index": "E.1/交付候选/100条交付索引.csv",
        "delivery_index_sha256": sha256(index_file) if index_file.exists() else None,
        "official_duration_statement_vs_measured": {
            "task_statement_s": "2.648–34.567",
            "measured_video_min_s": min(r["video_duration_s"] for r in stats),
            "measured_video_max_s": max(r["video_duration_s"] for r in stats),
            "handling": "保留原视频及标签，不按题面范围改写或筛除样本",
        },
        "failures": failures,
    }
    report_file = OUT / "提交前全量核验.json"
    report_file.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: report[k] for k in (
        "status", "sample_count", "raw_mp4_count", "raw_wav_count",
        "raw_audio_npz_count", "raw_visual_npz_count", "aligned_npz_count",
        "paper_appendix_sample_rows", "totals", "failures")}, ensure_ascii=False, indent=2))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
