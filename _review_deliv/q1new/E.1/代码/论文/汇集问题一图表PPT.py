"""Assemble the native workflow slide and seven paper figures for PPT review.

Slide 1 retains editable PowerPoint shapes. The remaining scientific plots are
high-resolution image objects; their plotting source is 重绘问题一图表_精修.py.
"""

from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[2]
PACKAGE = ROOT / "论文" / "图表精修"
FIGURES = PACKAGE / "论文插图"
SOURCE = PACKAGE / "流程图_可编辑" / "editable.pptx"
OUTPUT = PACKAGE / "问题一图表总览.pptx"
NAMES = [
    "fig02_evidence.png",
    "fig03_coverage.png",
    "fig04_distribution.png",
    "fig05_case_A.png",
    "fig06_case_B.png",
    "fig07_case_visual_missing.png",
    "fig08_case_asr_empty.png",
]


def main() -> None:
    prs = Presentation(SOURCE)
    blank = prs.slide_layouts[0]
    for name in NAMES:
        slide = prs.slides.add_slide(blank)
        for shape in list(slide.shapes):
            if shape.is_placeholder:
                shape._element.getparent().remove(shape._element)
        path = FIGURES / name
        from PIL import Image

        with Image.open(path) as im:
            width_px, height_px = im.size
        margin = int(min(prs.slide_width, prs.slide_height) * 0.025)
        area_w = prs.slide_width - 2 * margin
        area_h = prs.slide_height - 2 * margin
        scale = min(area_w / width_px, area_h / height_px)
        width = int(width_px * scale)
        height = int(height_px * scale)
        slide.shapes.add_picture(
            str(path),
            (prs.slide_width - width) // 2,
            (prs.slide_height - height) // 2,
            width=width,
            height=height,
        )
    prs.save(OUTPUT)
    print(f"{OUTPUT} | {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
